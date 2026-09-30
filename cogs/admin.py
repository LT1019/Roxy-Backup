import discord
from discord.ext import commands
from discord import app_commands
import aiosqlite
import asyncio
import io
import time
import psutil
from datetime import datetime, timedelta
from database import RoxyDatabase
from config import is_admin, is_admin_id


class OwnerOnlyView(discord.ui.View):
    """A view whose buttons and menus only respond to one user"""

    def __init__(self, owner_id: int, timeout: float = 300):
        super().__init__(timeout=timeout)
        self.owner_id = owner_id

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.owner_id:
            await interaction.response.send_message("❌ This menu isn't yours - run the command yourself to get your own.", ephemeral=True)
            return False
        return True


class ConfirmView(OwnerOnlyView):
    """Yes/No confirmation for destructive admin actions"""

    def __init__(self, owner_id: int):
        super().__init__(owner_id, timeout=30)
        self.confirmed = False

    @discord.ui.button(label='✅ Confirm', style=discord.ButtonStyle.danger)
    async def confirm(self, interaction, button):
        self.confirmed = True
        await interaction.response.edit_message(view=None)
        self.stop()

    @discord.ui.button(label='❌ Cancel', style=discord.ButtonStyle.secondary)
    async def cancel(self, interaction, button):
        await interaction.response.edit_message(content="❌ Cancelled.", embed=None, view=None)
        self.stop()


def build_excel_export(tables: dict, stats: dict) -> bytes:
    """Build an .xlsx workbook: a Summary sheet plus one sheet per database table"""
    import io
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="7B2FBE")  # Roxy purple

    def write_sheet(ws, columns, rows):
        ws.append(columns)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        id_columns = {c for c, column in enumerate(columns, start=1) if column == 'user_id'}
        for r, row in enumerate(rows, start=2):
            for c, value in enumerate(row, start=1):
                if c in id_columns and isinstance(value, int):
                    value = str(value)  # Discord IDs have 18 digits - Excel numbers only keep 15
                if isinstance(value, str):
                    value = ILLEGAL_CHARACTERS_RE.sub("", value)  # Control characters in names would break the file
                cell = ws.cell(row=r, column=c, value=value)
                if isinstance(value, str):
                    cell.data_type = 's'  # Always text - a name starting with "=" must never become a formula

        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for c, column in enumerate(columns, start=1):
            longest = max([len(str(column))] + [len(str(row[c - 1])) for row in rows[:500] if row[c - 1] is not None])
            ws.column_dimensions[get_column_letter(c)].width = min(longest + 2, 50)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [("Exported (UTC)", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S"))]
    summary_rows += [(key.replace('_', ' ').title(), value) for key, value in stats.items()]
    summary_rows += [(f"Rows: {name.replace('_', ' ').title()}", len(rows)) for name, (_, rows) in tables.items()]
    write_sheet(summary, ["Statistic", "Value"], summary_rows)

    for name, (columns, rows) in tables.items():
        write_sheet(workbook.create_sheet(name.replace('_', ' ').title()[:31]), columns, rows)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class DbStatsView(OwnerOnlyView):
    """!r dbstats buttons - only the admin can use them"""

    def __init__(self, cog, owner_id: int):
        super().__init__(owner_id, timeout=300)
        self.cog = cog
        self.message = None

    @discord.ui.button(label='📥 Export Excel', style=discord.ButtonStyle.success)
    async def export_excel(self, interaction, button):
        # Only the admin sees the file - it contains everyone's data
        await interaction.response.defer(ephemeral=True, thinking=True)
        try:
            tables = await self.cog.db.export_tables()
            stats = await self.cog.db.get_database_stats()
            data = await asyncio.to_thread(build_excel_export, tables, stats)
        except Exception as e:
            print(f"❌ Error building Excel export: {e}")
            await interaction.followup.send(f"❌ Couldn't build the export: {e}", ephemeral=True)
            return

        limit = interaction.guild.filesize_limit if interaction.guild else 10 * 1024 * 1024
        if len(data) > limit:
            await interaction.followup.send(
                f"❌ The export is {len(data) / 1024 / 1024:.1f} MB - over Discord's {limit / 1024 / 1024:.0f} MB upload limit. Use `!r backup` instead.",
                ephemeral=True
            )
            return

        filename = f"roxy_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
        await interaction.followup.send(
            content=f"📥 **Roxy database export** - {len(data) / 1024:,.0f} KB. Only you can see this file.",
            file=discord.File(io.BytesIO(data), filename=filename),
            ephemeral=True
        )
        print(f"👑 Admin {interaction.user} exported the database to Excel")

    async def on_timeout(self):
        if self.message:
            try:
                await self.message.edit(view=None)
            except discord.HTTPException:
                pass


async def redirect_to_slash(ctx, name: str) -> bool:
    """Private ("Only you can see this") replies only exist for slash commands.
    In a server, a text-command use is removed and pointed at the slash version. Returns True if redirected."""
    if ctx.interaction is not None or ctx.guild is None:
        return False  # Slash command, or a DM - already private
    try:
        await ctx.message.delete()
    except (discord.Forbidden, discord.NotFound):
        pass
    await ctx.send(f"🔒 Use **/{name}** instead - Discord only allows \"Only you can see this\" replies for slash commands.", delete_after=8)
    return True


async def confirm_action(ctx, prompt: str) -> bool:
    """Ask the admin to confirm a destructive action; True if confirmed"""
    view = ConfirmView(ctx.author.id)
    message = await ctx.send(f"⚠️ {prompt}", view=view)
    timed_out = await view.wait()
    if timed_out:
        await message.edit(content="⌛ Confirmation timed out.", view=None)
    return view.confirmed

class RoxyAdmin(commands.Cog):
    """Roxy's administrative commands and management system"""
    
    def __init__(self, bot):
        self.bot = bot
        self.db = RoxyDatabase()
    
    @commands.command(name='admin', aliases=['a'])
    @is_admin()
    async def admin_help(self, ctx):
        """Admin help menu - Only visible to admin, silent for others"""
        print(f"👑 Admin {ctx.author} accessed admin panel")
        
        embed = discord.Embed(
            title="👑 Roxy Admin Control Panel",
            description="**Admin-Only Commands** for managing Roxy Bot",
            color=discord.Color.red()
        )
        
        embed.add_field(
            name="👥 User Management",
            value="`!r addxp <@user> <amount>` - Give XP to user\n`!r setlevel <@user> <level>` - Set user level\n`!r resetuser <@user>` - Reset user stats\n`!r viewuser <@user>` - View detailed user data",
            inline=False
        )
        
        embed.add_field(
            name="🏆 Achievement Management",
            value="`!r ach` - Achievement control panel\n`!r ach users <achievement>` - Who has an achievement\n`!r giveach <@user> <achievement>` - Grant achievement\n`!r removeach <@user> <achievement>` - Remove achievement",
            inline=False
        )
        
        embed.add_field(
            name="🗄️ Database Management",
            value="`/dbstats` - Database statistics & Excel export (only you)\n`!r cleanup` - Clean inactive users\n`!r backup` - Create database backup\n`!r totalstats` - Global statistics\n`!r serverstats [server id]` - Server info & members",
            inline=False
        )
        
        embed.add_field(
            name="🤖 Bot Control",
            value="`!r setstatus <message>` - Set bot status (`clear` to resume rotation)\n`!r announce [#channel] <message>` - Send announcement\n`!r shutdown` - Shutdown bot\n`!r reload` - Reload cogs",
            inline=False
        )
        
        embed.add_field(
            name="🔧 Debug & Maintenance",
            value="`!r forceupdate <@user>` - Force update user\n`!r clearsessions` - End all active sessions\n`/logs [limit]` - View recent logs (only you can see them)\n`!r testxp` - Test XP system",
            inline=False
        )
        
        embed.add_field(
            name="⚡ Quick Access",
            value="`!admin` or `!a` - This admin panel\n`!addxp` - Quick XP commands work too",
            inline=False
        )
        
        embed.set_footer(text="👑 You are Roxy's Administrator | Admin commands are ignored for everyone else")
        
        await ctx.send(embed=embed)

    # Achievement Management Commands (unchanged from previous version)
    # The group itself has no check so that `!r ach list` stays public;
    # the control panel and the give/remove/users sub-commands are admin-only.
    @commands.group(name='ach', invoke_without_command=True)
    async def achievement_admin(self, ctx):
        """Admin achievement management control panel"""
        if not is_admin_id(ctx.author.id):
            return

        embed = discord.Embed(
            title="🏆 Admin Achievement Control Panel",
            description="**Manage user achievements and rewards**",
            color=discord.Color.gold()
        )
        
        embed.add_field(
            name="📋 **Available Commands**",
            value="`!r ach list` - View all available achievements\n`!r ach users <achievement>` - See who has an achievement\n`!r ach give <@user> <achievement>` - Grant achievement to user\n`!r ach remove <@user> <achievement>` - Remove achievement from user",
            inline=False
        )
        
        embed.add_field(
            name="🎯 **Built-in Achievement Categories**",
            value="• **Level Achievements** - Milestone levels (5, 10, 20, 50+)\n• **Gaming Achievements** - Playtime and session milestones\n• **Music Achievements** - Listening time and session milestones\n• **Social Achievements** - Message and activity milestones\n• **Special Achievements** - Custom and event achievements",
            inline=False
        )
        
        embed.add_field(
            name="💡 **Tips**",
            value="• Achievements are automatically awarded based on user activity\n• You can manually grant special achievements for events\n• Custom achievements are permanently stored\n• Use exact achievement names (case sensitive)",
            inline=False
        )
        
        embed.set_footer(text="👑 Achievement management • Use sub-commands for specific actions")
        await ctx.send(embed=embed)

    @achievement_admin.command(name='list')
    async def ach_list(self, ctx):
        """Enhanced paginated achievement list with dropdown navigation (public)"""
        
        # Define all achievement categories with expanded lists
        achievement_categories = {
            "Level Achievements": {
                "emoji": "⭐",
                "color": 0xf1c40f,
                "achievements": [
                    "`Active Member` - Reach Level 5",
                    "`Veteran` - Reach Level 10", 
                    "`Expert` - Reach Level 20",
                    "`Master` - Reach Level 30",
                    "`Legend` - Reach Level 50",
                    "`Elite` - Reach Level 100",
                    "`Champion` - Reach Level 200",
                    "`Grandmaster` - Reach Level 500",
                    "`Ascended` - Reach Level 750",
                    "`Eternal` - Reach Level 1000 (Highest title, levels continue infinitely)"
                ]
            },
            "Gaming Achievements": {
                "emoji": "🎮",
                "color": 0x9b59b6,
                "achievements": [
                    "`Gaming Started` - 1+ hour playtime",
                    "`Casual Player` - 5+ hours playtime",
                    "`Active Gamer` - 10+ hours playtime",
                    "`Dedicated Player` - 20+ hours playtime",
                    "`Gaming Master` - 50+ hours playtime",
                    "`Century Gamer` - 100+ hours playtime",
                    "`Gaming Legend` - 200+ hours playtime",
                    "`Gaming God` - 500+ hours playtime",
                    "`Gaming Deity` - 1000+ hours playtime",
                    "`Marathon Gamer` - 5+ hour single session"
                ]
            },
            "Music Achievements": {
                "emoji": "🎵",
                "color": 0x1db954,
                "achievements": [
                    "`First Listen` - 1+ hour listening time",
                    "`Getting Started` - 5+ hours listening time",
                    "`Music Fan` - 10+ hours listening time",
                    "`Active Listener` - 20+ hours listening time",
                    "`Dedicated Listener` - 50+ hours listening time",
                    "`Music Enthusiast` - 100+ hours listening time",
                    "`Music Master` - 200+ hours listening time",
                    "`Music Virtuoso` - 500+ hours listening time",
                    "`Music Deity` - 1000+ hours listening time",
                    "`Marathon Listener` - 8+ hour single session"
                ]
            },
            "Social Achievements": {
                "emoji": "💬",
                "color": 0x3498db,
                "achievements": [
                    "`Regular Chatter` - 50+ messages",
                    "`Conversationalist` - 100+ messages",
                    "`Chatterbox` - 500+ messages",
                    "`Social Butterfly` - 1000+ messages",
                    "`Community Voice` - 2500+ messages",
                    "`Server Legend` - 5000+ messages",
                    "`Message Master` - 10000+ messages",
                    "`Session Starter` - 20+ gaming sessions",
                    "`Session Expert` - 50+ gaming sessions",
                    "`Session Master` - 100+ gaming sessions"
                ]
            },
            "Special Achievements": {
                "emoji": "👑",
                "color": 0xe74c3c,
                "achievements": [
                    "`Roxy Bot Administrator` - Bot admin status",
                    "`Early Adopter` - Used Roxy before 2027 (UTC) - automatic",
                    "`Event Participant` - Participated in special events",
                    "`Community Helper` - Helped other users significantly"
                ]
            }
        }
        
        # Get category list for easier access
        category_names = list(achievement_categories.keys())
        current_category = 0  # Start with Level Achievements
        
        def create_achievement_embed(category_index):
            """Create embed for specific achievement category"""
            category_name = category_names[category_index]
            category_data = achievement_categories[category_name]
            
            embed = discord.Embed(
                title=f"🏆 {category_data['emoji']} {category_name}",
                description=f"**Complete list of {category_name.lower()} in Roxy Bot**",
                color=category_data['color']
            )
            
            # Add all achievements for this category
            achievements_text = "\n".join([f"• {achievement}" for achievement in category_data['achievements']])
            
            embed.add_field(
                name=f"📋 **All {category_name} ({len(category_data['achievements'])} total)**",
                value=achievements_text,
                inline=False
            )
            
            # Add category-specific tips
            tips = {
                "Level Achievements": "💡 **Tip:** Gain XP by chatting (5 XP/message), gaming (1 XP/minute), and listening (1 XP/2 minutes). Levels are infinite!",
                "Gaming Achievements": "💡 **Tip:** Set your Discord status to 'Playing [Game]' to track sessions",
                "Music Achievements": "💡 **Tip:** Listen to Spotify to track music sessions automatically",
                "Social Achievements": "💡 **Tip:** Stay active in chat and participate in gaming sessions",
                "Special Achievements": "💡 **Tip:** Early Adopter is automatic for anyone using Roxy before 2027 (UTC) - the rest are granted manually"
            }
            
            embed.add_field(
                name="💡 **How to Earn**",
                value=tips.get(category_name, "Stay active and engage with the community!"),
                inline=False
            )
            
            # Add navigation info
            embed.set_footer(
                text=f"🏆 Page {category_index + 1} of {len(category_names)} • Use dropdown or arrows to navigate"
            )
            
            return embed
        
        # Create dropdown select menu for category navigation
        class AchievementSelect(discord.ui.Select):
            def __init__(self):
                options = []
                for i, (category_name, category_data) in enumerate(achievement_categories.items()):
                    options.append(
                        discord.SelectOption(
                            label=category_name,
                            description=f"{len(category_data['achievements'])} achievements",
                            emoji=category_data['emoji'],
                            value=str(i)
                        )
                    )
                
                super().__init__(
                    placeholder="🏆 Select an achievement category...",
                    min_values=1,
                    max_values=1,
                    options=options,
                    row=0
                )
            
            async def callback(self, interaction):
                nonlocal current_category
                current_category = int(self.values[0])
                
                embed = create_achievement_embed(current_category)
                await interaction.response.edit_message(embed=embed, view=view)
        
        # Create view with dropdown and navigation buttons
        class AchievementView(OwnerOnlyView):
            def __init__(self):
                super().__init__(ctx.author.id, timeout=300)  # Only the person who opened it can navigate
                self.add_item(AchievementSelect())
            
            @discord.ui.button(label='◀️ Previous', style=discord.ButtonStyle.secondary, row=1)
            async def previous_category(self, interaction, button):
                nonlocal current_category
                current_category = (current_category - 1) % len(category_names)
                
                embed = create_achievement_embed(current_category)
                await interaction.response.edit_message(embed=embed, view=self)
            
            @discord.ui.button(label='▶️ Next', style=discord.ButtonStyle.secondary, row=1)
            async def next_category(self, interaction, button):
                nonlocal current_category
                current_category = (current_category + 1) % len(category_names)
                
                embed = create_achievement_embed(current_category)
                await interaction.response.edit_message(embed=embed, view=self)
            
            @discord.ui.button(label='🏠 Overview', style=discord.ButtonStyle.primary, row=1)
            async def show_overview(self, interaction, button):
                # Create overview embed showing all categories
                overview_embed = discord.Embed(
                    title="🏆 Achievement System Overview",
                    description="**Complete achievement system in Roxy Bot**",
                    color=discord.Color.gold()
                )
                
                for category_name, category_data in achievement_categories.items():
                    overview_embed.add_field(
                        name=f"{category_data['emoji']} **{category_name}**",
                        value=f"• **{len(category_data['achievements'])}** total achievements\n• Use dropdown to explore this category",
                        inline=True
                    )
                
                # Add summary statistics
                total_achievements = sum(len(cat_data['achievements']) for cat_data in achievement_categories.values())
                overview_embed.add_field(
                    name="📊 **System Summary**",
                    value=f"🏆 **{total_achievements}** total achievements available\n🎯 **{len(category_names)}** different categories\n💎 **Level 1000** \"Eternal\" is the highest title\n⚡ **Infinite** leveling system (no level cap)\n🎵 **Music tracking** via Spotify",
                    inline=False
                )
                
                overview_embed.set_footer(text="👑 Use the dropdown menu to explore specific achievement categories")
                
                await interaction.response.edit_message(embed=overview_embed, view=self)
            
            @discord.ui.button(label='❌ Close', style=discord.ButtonStyle.danger, row=1)
            async def close_menu(self, interaction, button):
                await interaction.response.edit_message(
                    embed=discord.Embed(
                        title="🏆 Achievement Menu Closed",
                        description="Use `!r ach list` to open again.",
                        color=discord.Color.red()
                    ),
                    view=None
                )
                self.stop()
            
            async def on_timeout(self):
                # Disable all components when timeout occurs
                for item in self.children:
                    try:
                        item.disabled = True  # type: ignore
                    except AttributeError:
                        pass  # Some items might not have disabled attribute
        
        # Create initial embed and view
        view = AchievementView()
        embed = create_achievement_embed(current_category)
        
        await ctx.send(embed=embed, view=view)

    @achievement_admin.command(name='users')
    @is_admin()
    async def ach_users(self, ctx, *, achievement: str):
        """See who has a granted achievement (Admin only)"""
        holders = await self.db.get_achievement_holders(achievement)

        embed = discord.Embed(
            title=f"🏆 {achievement}",
            color=discord.Color.gold()
        )
        if holders:
            lines = []
            for user_id, display_name, earned_date in holders:
                name = display_name or f"User {user_id}"
                earned = f"<t:{int(datetime.fromisoformat(earned_date).timestamp())}:d>" if earned_date else "unknown date"
                lines.append(f"• **{name}** - {earned}")
            embed.description = "\n".join(lines)[:4096]
            embed.set_footer(text=f"{len(holders)} user(s) have this achievement")
        else:
            embed.description = "Nobody has this achievement yet. Names are case sensitive."

        await ctx.send(embed=embed)

    @achievement_admin.command(name='give')
    @is_admin()
    async def ach_give(self, ctx, member: discord.Member, *, achievement: str):
        """Grant an achievement to a user (Admin only)"""
        await self.db.add_user(member.id, str(member), member.display_name)
        if await self.db.give_achievement(member.id, achievement):
            await ctx.send(f"🏆 Granted **{achievement}** to {member.mention}!")
        else:
            await ctx.send(f"ℹ️ {member.display_name} already has **{achievement}**.")

    @achievement_admin.command(name='remove')
    @is_admin()
    async def ach_remove(self, ctx, member: discord.Member, *, achievement: str):
        """Remove an achievement from a user (Admin only)"""
        if await self.db.remove_achievement(member.id, achievement):
            await ctx.send(f"🗑️ Removed **{achievement}** from {member.display_name}.")
        else:
            await ctx.send(f"ℹ️ {member.display_name} doesn't have **{achievement}** (names are case sensitive).")

    @commands.command(name='giveach')
    @is_admin()
    async def give_achievement_shortcut(self, ctx, member: discord.Member, *, achievement: str):
        """Shortcut for !r ach give"""
        await self.ach_give(ctx, member, achievement=achievement)

    @commands.command(name='removeach')
    @is_admin()
    async def remove_achievement_shortcut(self, ctx, member: discord.Member, *, achievement: str):
        """Shortcut for !r ach remove"""
        await self.ach_remove(ctx, member, achievement=achievement)

    # ==================== USER MANAGEMENT ====================

    @commands.command(name='addxp')
    @is_admin()
    async def add_xp(self, ctx, member: discord.Member, amount: int):
        """Give (or take, with a negative amount) XP (Admin only)"""
        await self.db.add_user(member.id, str(member), member.display_name)
        result = await self.db.add_xp(member.id, amount)
        if result is None:
            await ctx.send("❌ Couldn't update XP - check the console for details.")
            return

        new_xp, new_level = result
        verb = "Gave" if amount >= 0 else "Took"
        await ctx.send(f"✨ {verb} **{abs(amount):,} XP** {'to' if amount >= 0 else 'from'} {member.display_name} → **{new_xp:,} XP**, Level **{new_level}**")

    @commands.command(name='setlevel')
    @is_admin()
    async def set_level(self, ctx, member: discord.Member, level: int):
        """Set a user's level - XP is set to the minimum for that level (Admin only)"""
        if level < 1:
            await ctx.send("❌ Level must be 1 or higher.")
            return

        await self.db.add_user(member.id, str(member), member.display_name)
        xp = self.db.get_xp_for_level(level)
        await self.db.force_set_user_xp(member.id, xp, level)
        await ctx.send(f"⭐ Set {member.display_name} to Level **{level}** ({xp:,} XP)")

    @commands.command(name='resetuser')
    @is_admin()
    async def reset_user(self, ctx, member: discord.Member):
        """Reset all of a user's stats, sessions and achievements (Admin only)"""
        if not await confirm_action(ctx, f"Reset **all** stats, sessions and achievements for {member.display_name}? This can't be undone."):
            return

        await self.db.reset_user_stats(member.id)
        self.bot.active_sessions.pop(member.id, None)
        self.bot.active_listening.pop(member.id, None)
        await ctx.send(f"🔄 Reset all stats for {member.display_name}.")

    @commands.command(name='viewuser')
    @is_admin()
    async def view_user(self, ctx, member: discord.Member):
        """View detailed user data (Admin only)"""
        stats = await self.db.get_user_stats(member.id)
        if not stats:
            await ctx.send(f"❌ {member.display_name} isn't in the database yet.")
            return

        (user_id, username, display_name, join_date, total_messages, total_playtime,
         current_game, last_seen, level, xp, total_listening_time, current_song, current_artist) = stats[:13]
        session_stats = await self.db.get_session_statistics(member.id)
        listening_stats = await self.db.get_listening_statistics(member.id)
        custom_achievements = await self.db.get_user_custom_achievements(member.id)

        def fmt_time(seconds):
            return f"{seconds // 3600}h {(seconds % 3600) // 60}m"

        def fmt_date(iso):
            return f"<t:{int(datetime.fromisoformat(iso).timestamp())}:R>" if iso else "Never"

        embed = discord.Embed(
            title=f"🔍 User Data: {member.display_name}",
            color=discord.Color.red()
        )
        embed.set_thumbnail(url=member.display_avatar.url)

        embed.add_field(
            name="🪪 Identity",
            value=f"**ID:** {user_id}\n**Username:** {username}\n**First seen:** {fmt_date(join_date)}\n**Last seen:** {fmt_date(last_seen)}",
            inline=False
        )
        embed.add_field(
            name="📊 Progress",
            value=f"**Level:** {level}\n**XP:** {xp:,}\n**Next level at:** {self.db.get_xp_for_level(level + 1):,} XP\n**Messages:** {total_messages:,}",
            inline=True
        )
        embed.add_field(
            name="🎮 Gaming",
            value=f"**Total:** {fmt_time(total_playtime)}\n**Sessions:** {session_stats['total_sessions']}\n**Longest:** {session_stats['longest_session']}\n**Now:** {current_game or 'Nothing'}",
            inline=True
        )
        embed.add_field(
            name="🎵 Music",
            value=f"**Total:** {fmt_time(total_listening_time)}\n**Sessions:** {listening_stats['total_sessions']}\n**Longest:** {listening_stats['longest_session']}\n**Now:** {f'{current_song} by {current_artist}' if current_song else 'Nothing'}",
            inline=True
        )
        embed.add_field(
            name="🎖️ Granted Achievements",
            value="\n".join(f"• {name}" for name, _ in custom_achievements)[:1024] or "None",
            inline=False
        )

        await ctx.send(embed=embed)

    @commands.command(name='forceupdate')
    @is_admin()
    async def force_update(self, ctx, member: discord.Member):
        """Recalculate a user's totals and level from their sessions (Admin only)"""
        await self.db.add_user(member.id, str(member), member.display_name)
        total_playtime, total_listening_time = await self.db.force_refresh_user_stats(member.id)
        await ctx.send(
            f"🔄 Updated {member.display_name}: **{total_playtime // 3600}h {(total_playtime % 3600) // 60}m** gaming, "
            f"**{total_listening_time // 3600}h {(total_listening_time % 3600) // 60}m** listening"
        )

    # ==================== DATABASE MANAGEMENT ====================

    @commands.hybrid_command(name='dbstats', description="Database statistics and Excel export (admin only, only you can see it)")
    @app_commands.default_permissions(administrator=True)
    @is_admin()
    async def database_statistics(self, ctx):
        """Database statistics (Admin only)"""
        if await redirect_to_slash(ctx, 'dbstats'):
            return
        
        stats = await self.db.get_database_stats()
        if not stats:
            await ctx.send("❌ Couldn't load database statistics - check the console for details.")
            return

        import os
        db_size_kb = os.path.getsize(self.db.db_path) // 1024 if os.path.exists(self.db.db_path) else 0

        embed = discord.Embed(title="🗄️ Database Statistics", color=discord.Color.blue())
        embed.add_field(
            name="👥 Users",
            value=f"**Tracked:** {stats['total_users']:,}\n**Active (7d):** {stats['active_users']:,}\n**Highest level:** {stats['highest_level']}\n**Average level:** {stats['avg_level']:.1f}",
            inline=True
        )
        embed.add_field(
            name="📈 Activity",
            value=f"**Messages:** {stats['total_messages']:,}\n**Gaming:** {stats['total_playtime_hours']:,}h ({stats['total_sessions']:,} sessions)\n**Listening:** {stats['total_listening_hours']:,}h ({stats['total_listening_sessions']:,} sessions)",
            inline=True
        )
        embed.add_field(name="💾 File", value=f"`{self.db.db_path}` - {db_size_kb:,} KB", inline=False)
        embed.set_footer(text="👑 Export Excel sends the file privately - only you can see it")

        view = DbStatsView(self, ctx.author.id)
        view.message = await ctx.send(embed=embed, view=view, ephemeral=True)

    @commands.command(name='cleanup')
    @is_admin()
    async def cleanup_users(self, ctx, days: int = 30):
        """Remove users with no messages, gaming or listening who haven't been seen for X days (Admin only)"""
        if days < 1:
            await ctx.send("❌ Days must be 1 or more.")
            return

        if not await confirm_action(ctx, f"Delete users with **no activity at all** who haven't been seen in **{days} days**?"):
            return

        removed = await self.db.cleanup_inactive_users(days)
        await ctx.send(f"🧹 Removed **{removed}** inactive user(s).")

    @commands.command(name='backup')
    @is_admin()
    async def backup(self, ctx):
        """Create a database backup (Admin only)"""
        backup_path = await self.db.backup_database()
        if backup_path:
            await ctx.send(f"💾 Backup created: `{backup_path}`")
        else:
            await ctx.send("❌ Backup failed - check the console for details.")

    # ==================== BOT CONTROL ====================

    @commands.command(name='setstatus')
    @is_admin()
    async def set_status(self, ctx, *, message: str):
        """Set Roxy's status; `clear` resumes the rotating status (Admin only)"""
        if message.lower() == 'clear':
            self.bot.custom_status = None
            await self.bot.change_presence(activity=discord.Game(name="💜 Use !r help"))
            await ctx.send("✅ Custom status cleared - rotating status resumes.")
            return

        self.bot.custom_status = message[:128]
        await self.bot.change_presence(activity=discord.Game(name=self.bot.custom_status))
        await ctx.send(f"✅ Status set to **{self.bot.custom_status}** (use `!r setstatus clear` to resume rotation)")

    @commands.command(name='announce')
    @is_admin()
    async def announce(self, ctx, channel: discord.TextChannel = None, *, message: str):
        """Send an announcement embed, optionally to another channel (Admin only)"""
        target = channel or ctx.channel
        embed = discord.Embed(
            title="📢 Announcement",
            description=message,
            color=discord.Color.purple(),
            timestamp=datetime.now()
        )
        embed.set_footer(text=f"From {ctx.author.display_name}", icon_url=ctx.author.display_avatar.url)

        try:
            await target.send(embed=embed)
        except discord.Forbidden:
            await ctx.send(f"❌ I don't have permission to send messages in {target.mention}.")
            return

        if target != ctx.channel:
            await ctx.send(f"✅ Announcement sent to {target.mention}.")

    @commands.command(name='reload')
    @is_admin()
    async def reload_cogs(self, ctx):
        """Reload Roxy's cogs without restarting (Admin only)"""
        results = []
        for extension in ['cogs.stats', 'cogs.admin']:
            try:
                await self.bot.reload_extension(extension)
                results.append(f"✅ `{extension}`")
            except Exception as e:
                results.append(f"❌ `{extension}`: {e}")
        await ctx.send("🔄 Reload results:\n" + "\n".join(results))

    @commands.command(name='shutdown')
    @is_admin()
    async def shutdown(self, ctx):
        """Save active sessions and shut Roxy down (Admin only)"""
        if not await confirm_action(ctx, "Shut Roxy down? You'll need to start her again manually."):
            return

        await ctx.send("👋 Saving active sessions and shutting down...")
        await self.end_all_sessions()
        await self.bot.close()

    @commands.command(name='clearsessions')
    @is_admin()
    async def clear_sessions(self, ctx):
        """End and save every active session (Admin only)"""
        games, listens = await self.end_all_sessions()
        await ctx.send(
            f"🧹 Ended **{games}** gaming and **{listens}** listening session(s) - their time has been saved.\n"
            "Anyone still playing or listening will be picked up again when their status next changes."
        )

    async def end_all_sessions(self):
        """End every tracked session so its time and XP are credited"""
        games = list(self.bot.active_sessions)
        listens = list(self.bot.active_listening)
        self.bot.active_sessions.clear()
        self.bot.active_listening.clear()

        for user_id in games:
            await self.db.end_game_session(user_id)
        for user_id in listens:
            await self.db.end_listening_session(user_id)
        return len(games), len(listens)

    # Advanced Admin Commands
    @commands.command(name='totalstats')
    @is_admin()
    async def global_statistics(self, ctx):
        """Global statistics with a dropdown to view each section (Admin only)"""
        bot = self.bot
        db = self.db

        sections = {
            'all': {'label': 'All', 'emoji': '📊', 'description': 'Everything at a glance'},
            'reach': {'label': 'Global Reach', 'emoji': '🌍', 'description': 'Servers, users and live activity'},
            'activity': {'label': 'Activity Metrics', 'emoji': '📈', 'description': 'Messages, gaming and listening totals'},
            'progression': {'label': 'Progression', 'emoji': '🏆', 'description': 'Levels and XP'},
            'servers': {'label': 'Top Servers', 'emoji': '🏠', 'description': 'Biggest servers Roxy is in'},
            'performance': {'label': 'Bot Performance', 'emoji': '⚡', 'description': 'Uptime, CPU, RAM and latency'},
            'insights': {'label': 'Activity Insights', 'emoji': '🎮', 'description': 'Averages per user and session'},
            'achievements': {'label': 'Achievements', 'emoji': '🏅', 'description': 'Achievement system summary'},
        }
        state = {'section': 'all', 'server_page': 1, 'server_pages': 1}
        SERVERS_PER_PAGE = 10

        async def gather():
            """Collect fresh numbers every time the view changes"""
            db_stats = await db.get_database_stats()
            uptime = time.time() - bot.start_time if getattr(bot, 'start_time', 0) else 0.0
            memory = psutil.virtual_memory()
            return {
                'db': db_stats,
                'total_users': len(bot.users),
                'total_servers': len(bot.guilds),
                'active_gamers': len(getattr(bot, 'active_sessions', {})),
                'active_listeners': len(getattr(bot, 'active_listening', {})),
                'uptime': uptime,
                'cpu': psutil.cpu_percent(),
                'memory_mb': memory.used // 1024 // 1024,
                'memory_percent': memory.percent,
                'roxy_memory_mb': psutil.Process().memory_info().rss // 1024 // 1024,
                'latency': round(bot.latency * 1000) if bot.latency == bot.latency else 0,  # NaN while reconnecting
            }

        def add_reach(embed, d, detailed):
            s = d['db']
            embed.add_field(
                name="🌍 **Global Reach**",
                value=f"🏠 **{d['total_servers']}** servers\n👥 **{d['total_users']}** total users\n📊 **{s.get('total_users', 0)}** tracked users\n🎮 **{d['active_gamers']}** currently gaming\n🎵 **{d['active_listeners']}** currently listening",
                inline=not detailed
            )

        def add_activity(embed, d, detailed):
            s = d['db']
            embed.add_field(
                name="📈 **Activity Metrics**",
                value=f"💬 **{s.get('total_messages', 0):,}** total messages\n⏱️ **{s.get('total_playtime_hours', 0):,}h** total gaming\n🎵 **{s.get('total_listening_hours', 0):,}h** total listening\n🎲 **{s.get('total_sessions', 0):,}** gaming sessions\n🎶 **{s.get('total_listening_sessions', 0):,}** listening sessions\n📅 **{s.get('active_users', 0)}** active this week",
                inline=not detailed
            )

        def add_progression(embed, d, detailed):
            s = d['db']
            embed.add_field(
                name="🏆 **Progression Stats**",
                value=f"⭐ **Level {s.get('highest_level', 1)}** highest level\n📊 **{s.get('avg_level', 1.0):.1f}** average level\n✨ **{s.get('max_xp', 0):,}** highest XP\n🎯 **{s.get('max_xp', 0) // 5:,}** equivalent messages",
                inline=not detailed
            )

        def add_servers(embed, d, detailed):
            ranked = sorted(bot.guilds, key=lambda g: g.member_count or 0, reverse=True)
            if detailed:
                # Paged: 10 servers per page, Previous/Next buttons move through the rest
                state['server_pages'] = max(1, (len(ranked) + SERVERS_PER_PAGE - 1) // SERVERS_PER_PAGE)
                state['server_page'] = min(max(state['server_page'], 1), state['server_pages'])
                start = (state['server_page'] - 1) * SERVERS_PER_PAGE
                shown = ranked[start:start + SERVERS_PER_PAGE]
                title = f"🏠 **Top Servers {start + 1}-{start + len(shown)} of {len(ranked)}**"
            else:
                start = 0
                shown = ranked[:5]
                title = "🏠 **Top 5 Servers by Members**"

            server_info = "\n".join(
                f"**{i}.** {guild.name} - {guild.member_count} members" + (f" • `{guild.id}`" if detailed else "")
                for i, guild in enumerate(shown, start + 1)
            )
            embed.add_field(name=title, value=server_info or "No servers found", inline=False)
            if detailed:
                embed.add_field(name="💡 **Tip**", value="Use `!r serverstats <server id>` for full details on a server.", inline=False)

        def add_performance(embed, d, detailed):
            uptime_text = f"{int(d['uptime'] // 3600)}h {int((d['uptime'] % 3600) // 60)}m"
            value = f"🕐 **{uptime_text}** uptime\n🧠 **{d['memory_mb']}MB** RAM usage\n💻 **{d['cpu']:.1f}%** CPU usage\n📡 **{d['latency']}ms** latency"
            if detailed:
                value += f"\n📦 **{d['roxy_memory_mb']}MB** used by Roxy\n📊 **{d['memory_percent']:.1f}%** system RAM in use\n🧩 **{len(bot.cogs)}** cogs loaded"
            embed.add_field(name="⚡ **Bot Performance**", value=value, inline=not detailed)

        def add_insights(embed, d, detailed):
            s = d['db']
            avg_gaming_session_hours = s.get('total_playtime_hours', 0) / max(s.get('total_sessions', 1), 1)
            avg_listening_session_hours = s.get('total_listening_hours', 0) / max(s.get('total_listening_sessions', 1), 1)
            messages_per_user = s.get('total_messages', 0) / max(s.get('total_users', 1), 1)
            embed.add_field(
                name="🎮🎵 **Activity Insights**",
                value=f"⏱️ **{avg_gaming_session_hours:.1f}h** avg gaming session\n🎶 **{avg_listening_session_hours:.1f}h** avg listening session\n📝 **{messages_per_user:.0f}** avg messages per user\n🎯 **{(s.get('total_playtime_hours', 0) / max(d['total_users'], 1)):.1f}h** avg gaming per user\n🎵 **{(s.get('total_listening_hours', 0) / max(d['total_users'], 1)):.1f}h** avg listening per user\n📊 **{(s.get('active_users', 0) / max(s.get('total_users', 1), 1) * 100):.1f}%** weekly activity rate",
                inline=not detailed
            )

        def add_achievements(embed, d, detailed):
            embed.add_field(
                name="🏆 **Achievements & Milestones**",
                value="🎖️ **44** total achievement types\n🏅 **5** achievement categories\n🎮 **Gaming** tracking system\n🎵 **Music** listening tracking\n💎 **Level 1000** highest possible title\n⚡ **Infinite** leveling system",
                inline=not detailed
            )

        builders = {
            'reach': add_reach,
            'activity': add_activity,
            'progression': add_progression,
            'servers': add_servers,
            'performance': add_performance,
            'insights': add_insights,
            'achievements': add_achievements,
        }

        async def create_stats_embed():
            data = await gather()
            info = sections[state['section']]

            if state['section'] == 'all':
                embed = discord.Embed(
                    title="📊 Roxy Bot Global Statistics",
                    description="**Complete overview of Roxy Bot's performance and usage**",
                    color=discord.Color.gold()
                )
                for add_section in builders.values():
                    add_section(embed, data, detailed=False)
            else:
                embed = discord.Embed(
                    title=f"{info['emoji']} {info['label']}",
                    description=f"**{info['description']}**",
                    color=discord.Color.gold()
                )
                builders[state['section']](embed, data, detailed=True)

            embed.set_footer(
                text=f"👑 Generated by {ctx.author.display_name} • Use the dropdown to view each section",
                icon_url=ctx.author.display_avatar.url
            )
            if bot.user:
                embed.set_thumbnail(url=bot.user.display_avatar.url)
            return embed

        class StatsSelect(discord.ui.Select):
            def __init__(self):
                options = [
                    discord.SelectOption(label=info['label'], emoji=info['emoji'], description=info['description'],
                                         value=key, default=(key == state['section']))
                    for key, info in sections.items()
                ]
                super().__init__(placeholder="📊 Choose a statistics section...", min_values=1, max_values=1, options=options, row=0)

            async def callback(self, interaction):
                state['section'] = self.values[0]
                state['server_page'] = 1
                await self.view.refresh(interaction)

        class StatsView(OwnerOnlyView):
            def __init__(self):
                super().__init__(ctx.author.id, timeout=300)
                self.add_item(StatsSelect())

                # Page buttons only make sense on the Top Servers section
                if state['section'] == 'servers':
                    self.previous_page.disabled = state['server_page'] <= 1
                    self.next_page.disabled = state['server_page'] >= state['server_pages']
                else:
                    self.remove_item(self.previous_page)
                    self.remove_item(self.next_page)

            async def refresh(self, interaction):
                embed = await create_stats_embed()
                new_view = StatsView()
                new_view.message = interaction.message
                self.stop()  # Only the newest view stays alive
                await interaction.response.edit_message(embed=embed, view=new_view)

            @discord.ui.button(label='◀️ Previous', style=discord.ButtonStyle.secondary, row=1)
            async def previous_page(self, interaction, button):
                state['server_page'] -= 1
                await self.refresh(interaction)

            @discord.ui.button(label='▶️ Next', style=discord.ButtonStyle.secondary, row=1)
            async def next_page(self, interaction, button):
                state['server_page'] += 1
                await self.refresh(interaction)

            @discord.ui.button(label='🔄 Refresh', style=discord.ButtonStyle.primary, row=1)
            async def refresh_button(self, interaction, button):
                await self.refresh(interaction)

            @discord.ui.button(label='❌ Close', style=discord.ButtonStyle.danger, row=1)
            async def close_menu(self, interaction, button):
                self.stop()
                await interaction.response.edit_message(
                    embed=discord.Embed(
                        title="📊 Statistics Closed",
                        description="Use `!r totalstats` to open again.",
                        color=discord.Color.red()
                    ),
                    view=None
                )

            async def on_timeout(self):
                try:
                    await self.message.edit(view=None)
                except (AttributeError, discord.HTTPException):
                    pass

        try:
            embed = await create_stats_embed()
            view = StatsView()
            view.message = await ctx.send(embed=embed, view=view)
        except Exception as e:
            await ctx.send(f"❌ Error generating global statistics: {e}")
            print(f"❌ Error in totalstats: {e}")

    @commands.command(name='serverstats', aliases=['serverinfo'])
    @is_admin()
    async def server_statistics(self, ctx, server_id: int = None):
        """Full information and member list for a server Roxy is in (Admin only)"""
        if server_id is None:
            if ctx.guild is None:
                await ctx.send("❌ Use this in a server, or give a server ID: `!r serverstats <server id>`")
                return
            guild = ctx.guild
        else:
            guild = self.bot.get_guild(server_id)
            if guild is None:
                await ctx.send(f"❌ Roxy isn't in a server with ID `{server_id}`. Use `!r totalstats` → Top Servers to see server IDs.")
                return

        bot = self.bot
        db = self.db
        MEMBERS_PER_PAGE = 15
        sections = {
            'overview': {'label': 'Overview', 'emoji': '🏠', 'description': 'Server information'},
            'members': {'label': 'Members', 'emoji': '👥', 'description': 'Everyone who has joined, oldest first'},
            'roxy': {'label': 'Roxy Stats', 'emoji': '📊', 'description': "This server's activity tracked by Roxy"},
        }
        state = {'section': 'overview', 'page': 1, 'pages': 1}

        def fmt_date(dt):
            return f"<t:{int(dt.timestamp())}:D> (<t:{int(dt.timestamp())}:R>)" if dt else "Unknown"

        def create_overview_embed():
            humans = sum(1 for m in guild.members if not m.bot)
            bots = sum(1 for m in guild.members if m.bot)
            online = sum(1 for m in guild.members if m.status != discord.Status.offline)

            embed = discord.Embed(title=f"🏠 {guild.name}", description=guild.description or None, color=discord.Color.blurple())
            if guild.icon:
                embed.set_thumbnail(url=guild.icon.url)
            if guild.banner:
                embed.set_image(url=guild.banner.url)

            embed.add_field(
                name="🪪 **General**",
                value=f"**ID:** `{guild.id}`\n**Owner:** {guild.owner.mention if guild.owner else 'Unknown'} (`{guild.owner_id}`)\n**Created:** {fmt_date(guild.created_at)}\n**Roxy joined:** {fmt_date(guild.me.joined_at if guild.me else None)}",
                inline=False
            )
            embed.add_field(
                name="👥 **Members**",
                value=f"**Total:** {guild.member_count:,}\n**Humans:** {humans:,}\n**Bots:** {bots:,}\n**Online:** {online:,}",
                inline=True
            )
            embed.add_field(
                name="💬 **Channels**",
                value=f"**Text:** {len(guild.text_channels)}\n**Voice:** {len(guild.voice_channels)}\n**Categories:** {len(guild.categories)}\n**Forums:** {len(guild.forums)}\n**Stage:** {len(guild.stage_channels)}",
                inline=True
            )
            embed.add_field(
                name="✨ **Extras**",
                value=f"**Roles:** {len(guild.roles) - 1}\n**Emojis:** {len(guild.emojis)}/{guild.emoji_limit}\n**Stickers:** {len(guild.stickers)}/{guild.sticker_limit}\n**Boost tier:** {guild.premium_tier}\n**Boosts:** {guild.premium_subscription_count or 0}",
                inline=True
            )
            embed.add_field(
                name="🔒 **Settings**",
                value=f"**Verification:** {str(guild.verification_level).title()}\n**Content filter:** {str(guild.explicit_content_filter).replace('_', ' ').title()}\n**2FA for mods:** {'Yes' if guild.mfa_level else 'No'}\n**Locale:** {guild.preferred_locale}",
                inline=True
            )
            if guild.features:
                embed.add_field(
                    name="🎁 **Features**",
                    value=", ".join(f.replace('_', ' ').title() for f in sorted(guild.features))[:1024],
                    inline=False
                )
            return embed

        def create_members_embed():
            members = sorted(guild.members, key=lambda m: m.joined_at or discord.utils.utcnow())
            state['pages'] = max(1, (len(members) + MEMBERS_PER_PAGE - 1) // MEMBERS_PER_PAGE)
            state['page'] = min(max(state['page'], 1), state['pages'])
            start = (state['page'] - 1) * MEMBERS_PER_PAGE

            lines = []
            for number, member in enumerate(members[start:start + MEMBERS_PER_PAGE], start + 1):
                joined = f"<t:{int(member.joined_at.timestamp())}:R>" if member.joined_at else "unknown"
                tags = (" 🤖" if member.bot else "") + (" 👑" if member.id == guild.owner_id else "")
                lines.append(f"`#{number}` **{member.display_name}**{tags} (@{member.name}) • joined {joined}")

            embed = discord.Embed(
                title=f"👥 Members of {guild.name}",
                description="\n".join(lines) or "No members found",
                color=discord.Color.blurple()
            )
            embed.add_field(
                name="📊 **Showing**",
                value=f"{start + 1}-{start + len(lines)} of {len(members):,} members (oldest first)",
                inline=False
            )
            return embed

        async def create_roxy_embed():
            member_ids = {m.id for m in guild.members if not m.bot}
            tracked = [row for row in await db.get_all_users_admin(limit=1000000) if row[0] in member_ids]
            # rows: user_id, display_name, total_messages, total_playtime, total_listening_time, level, xp, last_seen
            total_messages = sum(r[2] or 0 for r in tracked)
            total_playtime = sum(r[3] or 0 for r in tracked)
            total_listening = sum(r[4] or 0 for r in tracked)
            week_ago = (datetime.now() - timedelta(days=7)).isoformat()
            active_week = sum(1 for r in tracked if r[7] and r[7] > week_ago)
            gaming_now = sum(1 for uid in getattr(bot, 'active_sessions', {}) if uid in member_ids)
            listening_now = sum(1 for uid in getattr(bot, 'active_listening', {}) if uid in member_ids)

            embed = discord.Embed(title=f"📊 Roxy Stats for {guild.name}", color=discord.Color.gold())
            embed.add_field(
                name="👥 **Tracking**",
                value=f"**Tracked members:** {len(tracked):,} of {len(member_ids):,} humans\n**Active this week:** {active_week:,}\n**Gaming now:** {gaming_now}\n**Listening now:** {listening_now}",
                inline=False
            )
            embed.add_field(
                name="📈 **Totals**",
                value=f"💬 **{total_messages:,}** messages\n🎮 **{total_playtime // 3600:,}h** gaming\n🎵 **{total_listening // 3600:,}h** listening",
                inline=True
            )
            top_level = max(tracked, key=lambda r: r[6] or 0, default=None)
            top_chatter = max(tracked, key=lambda r: r[2] or 0, default=None)
            embed.add_field(
                name="🏆 **Top Members**",
                value=(f"⭐ **Highest level:** {top_level[1]} (Lv.{top_level[5]})\n💬 **Most messages:** {top_chatter[1]} ({top_chatter[2]:,})"
                       if tracked else "No tracked members yet"),
                inline=True
            )
            return embed

        async def create_embed():
            if state['section'] == 'overview':
                embed = create_overview_embed()
            elif state['section'] == 'members':
                embed = create_members_embed()
            else:
                embed = await create_roxy_embed()

            page_text = f"Page {state['page']}/{state['pages']} • " if state['section'] == 'members' else ""
            embed.set_footer(text=f"👑 {page_text}Server ID {guild.id} • Use the dropdown to switch views")
            return embed

        class ServerSelect(discord.ui.Select):
            def __init__(self):
                options = [
                    discord.SelectOption(label=info['label'], emoji=info['emoji'], description=info['description'],
                                         value=key, default=(key == state['section']))
                    for key, info in sections.items()
                ]
                super().__init__(placeholder="🏠 Choose a view...", min_values=1, max_values=1, options=options, row=0)

            async def callback(self, interaction):
                state['section'] = self.values[0]
                state['page'] = 1
                await self.view.refresh(interaction)

        class ServerView(OwnerOnlyView):
            def __init__(self):
                super().__init__(ctx.author.id, timeout=300)
                self.add_item(ServerSelect())

                # Page buttons only on the member list
                if state['section'] == 'members':
                    self.previous_page.disabled = state['page'] <= 1
                    self.next_page.disabled = state['page'] >= state['pages']
                else:
                    self.remove_item(self.previous_page)
                    self.remove_item(self.next_page)

            async def refresh(self, interaction):
                embed = await create_embed()
                new_view = ServerView()
                new_view.message = interaction.message
                self.stop()  # Only the newest view stays alive
                await interaction.response.edit_message(embed=embed, view=new_view)

            @discord.ui.button(label='◀️ Previous', style=discord.ButtonStyle.secondary, row=1)
            async def previous_page(self, interaction, button):
                state['page'] -= 1
                await self.refresh(interaction)

            @discord.ui.button(label='▶️ Next', style=discord.ButtonStyle.secondary, row=1)
            async def next_page(self, interaction, button):
                state['page'] += 1
                await self.refresh(interaction)

            @discord.ui.button(label='🔄 Refresh', style=discord.ButtonStyle.primary, row=1)
            async def refresh_button(self, interaction, button):
                await self.refresh(interaction)

            @discord.ui.button(label='❌ Close', style=discord.ButtonStyle.danger, row=1)
            async def close_menu(self, interaction, button):
                self.stop()
                await interaction.response.edit_message(
                    embed=discord.Embed(
                        title="🏠 Server Stats Closed",
                        description="Use `!r serverstats [server id]` to open again.",
                        color=discord.Color.red()
                    ),
                    view=None
                )

            async def on_timeout(self):
                try:
                    await self.message.edit(view=None)
                except (AttributeError, discord.HTTPException):
                    pass

        embed = await create_embed()
        view = ServerView()
        view.message = await ctx.send(embed=embed, view=view)

    @commands.hybrid_command(name='logs', description="Recent activity logs (admin only, only you can see it)")
    @app_commands.default_permissions(administrator=True)
    @app_commands.describe(limit="Entries per page (1-50, default 10)")
    @is_admin()
    async def view_recent_logs(self, ctx, limit: int = 10):
        """Enhanced paginated logs with dropdown navigation (Admin only)"""
        if await redirect_to_slash(ctx, 'logs'):
            return
        
        # Enforce limit constraints: default 10, max 50
        if limit < 1:
            limit = 10
        elif limit > 50:
            limit = 50
        
        # Define all log categories
        log_categories = {
            "Recent User Activity": {
                "emoji": "👥",
                "color": 0x3498db,
                "description": "Latest user interactions and activity"
            },
            "Recent Gaming Sessions": {
                "emoji": "🎮", 
                "color": 0x9b59b6,
                "description": "Recently completed gaming sessions"
            },
            "Recent Listening Sessions": {
                "emoji": "🎵",
                "color": 0x1db954,
                "description": "Recently completed music listening sessions"
            },
            "Currently Active Gaming": {
                "emoji": "🔴",
                "color": 0xe74c3c,
                "description": "Live gaming sessions in progress"
            },
            "Currently Active Listening": {
                "emoji": "🟢",
                "color": 0x27ae60,
                "description": "Live music listening sessions in progress"
            },
            "Top Recent Levels": {
                "emoji": "⭐",
                "color": 0xf1c40f,
                "description": "Recent level achievements and top users"
            }
        }
        
        # Get category list for easier access
        category_names = list(log_categories.keys())
        current_category = 0  # Start with Recent User Activity
        current_page = 1  # Add page tracking
        
        async def create_user_activity_embed(page=1):
            """Create Recent User Activity embed with pagination"""
            embed = discord.Embed(
                title="👥 Recent User Activity",
                description="**Latest user interactions and activity**",
                color=0x3498db
            )
            
            # Calculate offset for pagination
            offset = (page - 1) * limit
            
            # Get recent user activities with offset
            recent_users = []
            total_count = 0
            try:
                async with aiosqlite.connect(self.db.db_path) as db:
                    # First get total count
                    async with db.execute("SELECT COUNT(*) FROM users") as cursor:
                        result = await cursor.fetchone()
                        total_count = result[0] if result else 0
                    
                    # Then get paginated results
                    async with db.execute("""
                        SELECT username, last_seen, total_messages, total_listening_time, level, xp
                        FROM users 
                        ORDER BY last_seen DESC 
                        LIMIT ? OFFSET ?
                    """, (limit, offset)) as cursor:
                        results = await cursor.fetchall()
                        
                    for username, last_seen, messages, listening_time, level, xp in results:
                        if last_seen:
                            last_seen_dt = datetime.fromisoformat(last_seen)
                            time_ago = f"<t:{int(last_seen_dt.timestamp())}:R>"
                            short_name = username[:15] + "..." if len(username) > 15 else username
                            listening_hours = listening_time // 3600
                            recent_users.append(f"• **{short_name}** (Lv.{level}, {messages:,} msgs, {listening_hours}h music) - {time_ago}")
            except Exception as e:
                recent_users = [f"❌ Error loading user data: {e}"]
            
            # Calculate total pages
            total_pages = (total_count + limit - 1) // limit if total_count > 0 else 1
            
            # Add users field
            user_text = "\n".join(recent_users) if recent_users else "No recent activity"
            embed.add_field(
                name=f"📊 **Recent Active Users (Page {page}/{total_pages})**",
                value=user_text[:1024],  # Discord field limit
                inline=False
            )
            
            # Add summary stats with music listening
            try:
                async with aiosqlite.connect(self.db.db_path) as db:
                    # Users active in last 24 hours
                    yesterday = (datetime.now().timestamp() - 86400)
                    async with db.execute("""
                        SELECT COUNT(*) FROM users 
                        WHERE datetime(last_seen) > datetime(?, 'unixepoch')
                    """, (yesterday,)) as cursor:
                        result = await cursor.fetchone()
                        active_24h = result[0] if result else 0
                    
                    # Users active this week
                    week_ago = (datetime.now().timestamp() - 604800)
                    async with db.execute("""
                        SELECT COUNT(*) FROM users 
                        WHERE datetime(last_seen) > datetime(?, 'unixepoch')
                    """, (week_ago,)) as cursor:
                        result = await cursor.fetchone()
                        active_week = result[0] if result else 0
                    
                    # Total listening hours
                    async with db.execute("""
                        SELECT SUM(total_listening_time) FROM users
                    """) as cursor:
                        result = await cursor.fetchone()
                        total_listening_seconds = result[0] if result and result[0] else 0
                        total_listening_hours = total_listening_seconds // 3600
                
                embed.add_field(
                    name="📈 **Activity Summary**",
                    value=f"🕐 **{active_24h}** users active (24h)\n📅 **{active_week}** users active (7d)\n👥 **{total_count}** total users tracked\n🎵 **{total_listening_hours:,}h** total music listened\n⚡ **{len(recent_users)}** users on this page",
                    inline=False
                )
            except Exception as e:
                embed.add_field(
                    name="📈 **Activity Summary**", 
                    value=f"❌ Error loading summary: {e}",
                    inline=False
                )
            
            return embed, total_pages
        
        async def create_listening_sessions_embed(page=1):
            """Create Recent Listening Sessions embed with pagination"""
            embed = discord.Embed(
                title="🎵 Recent Listening Sessions",
                description="**Recently completed music listening sessions**",
                color=0x1db954
            )
            
            # Calculate offset for pagination
            offset = (page - 1) * limit
            
            # Get recent listening sessions
            recent_sessions = []
            total_count = 0
            try:
                async with aiosqlite.connect(self.db.db_path) as db:
                    # Get total count
                    async with db.execute("""
                        SELECT COUNT(*) FROM listening_sessions 
                        WHERE end_time IS NOT NULL
                    """) as cursor:
                        result = await cursor.fetchone()
                        total_count = result[0] if result else 0
                    
                    # Get paginated results
                    async with db.execute("""
                        SELECT users.username, listening_sessions.song_title, listening_sessions.artist_name, listening_sessions.end_time, listening_sessions.duration
                        FROM listening_sessions 
                        JOIN users ON listening_sessions.user_id = users.user_id
                        WHERE listening_sessions.end_time IS NOT NULL
                        ORDER BY listening_sessions.end_time DESC 
                        LIMIT ? OFFSET ?
                    """, (limit, offset)) as cursor:
                        results = await cursor.fetchall()
                        
                    for username, song_title, artist_name, end_time, duration in results:
                        if end_time and duration:
                            end_dt = datetime.fromisoformat(end_time)
                            time_ago = f"<t:{int(end_dt.timestamp())}:R>"
                            duration_mins = duration // 60
                            duration_secs = duration % 60
                            
                            short_name = username[:12] + "..." if len(username) > 12 else username
                            short_song = song_title[:20] + "..." if len(song_title) > 20 else song_title
                            short_artist = artist_name[:15] + "..." if len(artist_name) > 15 else artist_name
                            
                            recent_sessions.append(f"• **{short_name}** listened to *{short_song}* by *{short_artist}* ({duration_mins}m {duration_secs}s) - {time_ago}")
            except Exception as e:
                recent_sessions = [f"❌ Error loading listening data: {e}"]
            
            # Calculate total pages
            total_pages = (total_count + limit - 1) // limit if total_count > 0 else 1
            
            sessions_text = "\n".join(recent_sessions) if recent_sessions else "No recent sessions"
            
            embed.add_field(
                name=f"🎶 **Recent Listening Sessions (Page {page}/{total_pages})**",
                value=sessions_text[:1024],
                inline=False
            )
            
            # Add listening summary
            try:
                async with aiosqlite.connect(self.db.db_path) as db:
                    # Sessions in last 24 hours
                    yesterday = (datetime.now().timestamp() - 86400)
                    async with db.execute("""
                        SELECT COUNT(*), SUM(duration) FROM listening_sessions 
                        WHERE datetime(end_time) > datetime(?, 'unixepoch') AND duration IS NOT NULL
                    """, (yesterday,)) as cursor:
                        result = await cursor.fetchone()
                        sessions_24h = 0
                        listening_time_24h = 0
                        if result:
                            sessions_24h = result[0] if result[0] else 0
                            listening_time_24h = (result[1] // 3600) if result[1] else 0
                    
                    # Total sessions today
                    today = datetime.now().date().isoformat()
                    async with db.execute("""
                        SELECT COUNT(*) FROM listening_sessions 
                        WHERE date(end_time) = ? AND duration IS NOT NULL
                    """, (today,)) as cursor:
                        result = await cursor.fetchone()
                        sessions_today = result[0] if result else 0
                
                embed.add_field(
                    name="📊 **Listening Summary**",
                    value=f"🕐 **{sessions_24h}** sessions (24h)\n⏱️ **{listening_time_24h}h** listening (24h)\n📅 **{sessions_today}** sessions today\n🎵 **{len(recent_sessions)}** sessions on this page",
                    inline=False
                )
            except Exception as e:
                embed.add_field(
                    name="📊 **Listening Summary**",
                    value=f"❌ Error loading summary: {e}",
                    inline=False
                )
            
            return embed, total_pages
        
        async def create_active_listening_embed(page=1):
            """Create Currently Active Listening embed (no pagination needed)"""
            embed = discord.Embed(
                title="🟢 Currently Active Listening",
                description="**Live music listening sessions in progress**",
                color=0x27ae60
            )
            
            # Get active listening sessions
            active_sessions_info = []
            session_details = []
            
            active_listening = getattr(self.bot, 'active_listening', {})
            
            for user_id, music_info in active_listening.items():
                user = self.bot.get_user(user_id)
                if user:
                    short_name = user.name[:12] + "..." if len(user.name) > 12 else user.name
                    song = music_info.get('song', 'Unknown')[:25] + "..." if len(music_info.get('song', 'Unknown')) > 25 else music_info.get('song', 'Unknown')
                    artist = music_info.get('artist', 'Unknown')[:20] + "..." if len(music_info.get('artist', 'Unknown')) > 20 else music_info.get('artist', 'Unknown')
                    
                    active_sessions_info.append(f"• **{short_name}** - *{song}* by *{artist}*")
                    
                    # Try to get session start time
                    try:
                        async with aiosqlite.connect(self.db.db_path) as db:
                            async with db.execute("""
                                SELECT start_time FROM listening_sessions 
                                WHERE user_id = ? AND end_time IS NULL 
                                ORDER BY start_time DESC LIMIT 1
                            """, (user_id,)) as cursor:
                                result = await cursor.fetchone()
                                if result:
                                    start_dt = datetime.fromisoformat(result[0])
                                    duration = datetime.now() - start_dt
                                    duration_mins = int(duration.total_seconds() // 60)
                                    session_details.append(f"• **{short_name}** - *{song}* by *{artist}* ({duration_mins}m)")
                                else:
                                    session_details.append(f"• **{short_name}** - *{song}* by *{artist}* (just started)")
                    except:
                        session_details.append(f"• **{short_name}** - *{song}* by *{artist}* (duration unknown)")
            
            active_text = "\n".join(session_details[:15]) if session_details else "No active listening sessions"
            
            embed.add_field(
                name=f"🎵 **Active Sessions ({len(active_sessions_info)})**",
                value=active_text[:1024],
                inline=False
            )
            
            # Add live statistics
            total_active = len(active_listening)
            total_users = len(self.bot.users)
            activity_rate = (total_active / max(total_users, 1)) * 100
            
            # Get popular artists currently being listened to
            artist_counts = {}
            for music_info in active_listening.values():
                artist = music_info.get('artist', 'Unknown')
                short_artist = artist[:20] + "..." if len(artist) > 20 else artist
                artist_counts[short_artist] = artist_counts.get(short_artist, 0) + 1
            
            popular_artists = sorted(artist_counts.items(), key=lambda x: x[1], reverse=True)[:5]
            popular_artists_text = "\n".join([f"• **{artist}** ({count} listeners)" for artist, count in popular_artists]) if popular_artists else "No artists currently active"
            
            embed.add_field(
                name="📊 **Live Statistics**",
                value=f"🎵 **{total_active}** active listeners\n👥 **{total_users}** total users\n📈 **{activity_rate:.1f}%** listening rate\n⚡ **Real-time** tracking",
                inline=True
            )
            
            embed.add_field(
                name="🏆 **Popular Artists Now**",
                value=popular_artists_text[:1024],
                inline=True
            )
            
            return embed, 1  # No pagination for active listening
        
        async def create_gaming_sessions_embed(page=1):
            """Create Recent Gaming Sessions embed with pagination"""
            embed = discord.Embed(
                title="🎮 Recent Gaming Sessions",
                description="**Recently completed gaming sessions**",
                color=0x9b59b6
            )

            offset = (page - 1) * limit
            recent_sessions = []
            total_count = 0
            try:
                async with aiosqlite.connect(self.db.db_path) as db:
                    async with db.execute("""
                        SELECT COUNT(*) FROM game_sessions WHERE end_time IS NOT NULL
                    """) as cursor:
                        result = await cursor.fetchone()
                        total_count = result[0] if result else 0

                    async with db.execute("""
                        SELECT users.username, game_sessions.game_name, game_sessions.end_time, game_sessions.duration
                        FROM game_sessions
                        JOIN users ON game_sessions.user_id = users.user_id
                        WHERE game_sessions.end_time IS NOT NULL
                        ORDER BY game_sessions.end_time DESC
                        LIMIT ? OFFSET ?
                    """, (limit, offset)) as cursor:
                        results = await cursor.fetchall()

                    for username, game_name, end_time, duration in results:
                        if end_time and duration:
                            time_ago = f"<t:{int(datetime.fromisoformat(end_time).timestamp())}:R>"
                            short_name = username[:12] + "..." if len(username) > 12 else username
                            short_game = game_name[:25] + "..." if len(game_name) > 25 else game_name
                            recent_sessions.append(f"• **{short_name}** played *{short_game}* ({duration // 3600}h {(duration % 3600) // 60}m) - {time_ago}")
            except Exception as e:
                recent_sessions = [f"❌ Error loading gaming data: {e}"]

            total_pages = (total_count + limit - 1) // limit if total_count > 0 else 1

            embed.add_field(
                name=f"🕹️ **Recent Gaming Sessions (Page {page}/{total_pages})**",
                value=("\n".join(recent_sessions) or "No recent sessions")[:1024],
                inline=False
            )

            try:
                async with aiosqlite.connect(self.db.db_path) as db:
                    yesterday = (datetime.now().timestamp() - 86400)
                    async with db.execute("""
                        SELECT COUNT(*), SUM(duration) FROM game_sessions
                        WHERE datetime(end_time) > datetime(?, 'unixepoch') AND duration IS NOT NULL
                    """, (yesterday,)) as cursor:
                        result = await cursor.fetchone()
                        sessions_24h = result[0] if result and result[0] else 0
                        gaming_hours_24h = (result[1] // 3600) if result and result[1] else 0

                    async with db.execute("""
                        SELECT game_name, COUNT(*) FROM game_sessions
                        WHERE datetime(end_time) > datetime(?, 'unixepoch') AND duration IS NOT NULL
                        GROUP BY game_name ORDER BY COUNT(*) DESC LIMIT 1
                    """, (yesterday,)) as cursor:
                        top_game = await cursor.fetchone()

                top_game_text = f"{top_game[0]} ({top_game[1]} sessions)" if top_game else "None"
                embed.add_field(
                    name="📊 **Gaming Summary**",
                    value=f"🕐 **{sessions_24h}** sessions (24h)\n⏱️ **{gaming_hours_24h}h** played (24h)\n🔥 **Most played (24h):** {top_game_text}",
                    inline=False
                )
            except Exception as e:
                embed.add_field(name="📊 **Gaming Summary**", value=f"❌ Error loading summary: {e}", inline=False)

            return embed, total_pages

        async def create_active_gaming_embed(page=1):
            """Create Currently Active Gaming embed (no pagination needed)"""
            embed = discord.Embed(
                title="🔴 Currently Active Gaming",
                description="**Live gaming sessions in progress**",
                color=0xe74c3c
            )

            active_sessions = getattr(self.bot, 'active_sessions', {})
            session_details = []

            try:
                async with aiosqlite.connect(self.db.db_path) as db:
                    for user_id, game in active_sessions.items():
                        user = self.bot.get_user(user_id)
                        if not user:
                            continue
                        short_name = user.name[:12] + "..." if len(user.name) > 12 else user.name
                        short_game = game[:25] + "..." if len(game) > 25 else game

                        async with db.execute("""
                            SELECT start_time FROM game_sessions
                            WHERE user_id = ? AND end_time IS NULL
                            ORDER BY start_time DESC LIMIT 1
                        """, (user_id,)) as cursor:
                            result = await cursor.fetchone()

                        if result:
                            minutes = int((datetime.now() - datetime.fromisoformat(result[0])).total_seconds() // 60)
                            session_details.append(f"• **{short_name}** - *{short_game}* ({minutes}m)")
                        else:
                            session_details.append(f"• **{short_name}** - *{short_game}* (just started)")
            except Exception as e:
                session_details = [f"❌ Error loading session times: {e}"]

            embed.add_field(
                name=f"🎮 **Active Sessions ({len(active_sessions)})**",
                value=("\n".join(session_details[:15]) or "No active gaming sessions")[:1024],
                inline=False
            )

            # Most popular games right now
            game_counts = {}
            for game in active_sessions.values():
                game_counts[game] = game_counts.get(game, 0) + 1
            popular_games = sorted(game_counts.items(), key=lambda x: x[1], reverse=True)[:5]

            embed.add_field(
                name="📊 **Live Statistics**",
                value=f"🎮 **{len(active_sessions)}** active gamers\n👥 **{len(self.bot.users)}** total users\n📈 **{len(active_sessions) / max(len(self.bot.users), 1) * 100:.1f}%** gaming rate",
                inline=True
            )
            embed.add_field(
                name="🏆 **Popular Games Now**",
                value=("\n".join(f"• **{game[:25]}** ({count} playing)" for game, count in popular_games) or "No games currently active")[:1024],
                inline=True
            )

            return embed, 1

        async def create_top_levels_embed(page=1):
            """Create Top Recent Levels embed with pagination"""
            embed = discord.Embed(
                title="⭐ Top Recent Levels",
                description="**Highest level users and their recent activity**",
                color=0xf1c40f
            )

            offset = (page - 1) * limit
            entries = []
            total_count = 0
            try:
                async with aiosqlite.connect(self.db.db_path) as db:
                    async with db.execute("SELECT COUNT(*) FROM users WHERE xp > 0") as cursor:
                        result = await cursor.fetchone()
                        total_count = result[0] if result else 0

                    async with db.execute("""
                        SELECT username, level, xp, last_seen FROM users
                        WHERE xp > 0
                        ORDER BY xp DESC
                        LIMIT ? OFFSET ?
                    """, (limit, offset)) as cursor:
                        results = await cursor.fetchall()

                    for rank, (username, level, xp, last_seen) in enumerate(results, start=offset + 1):
                        short_name = username[:15] + "..." if len(username) > 15 else username
                        seen = f"<t:{int(datetime.fromisoformat(last_seen).timestamp())}:R>" if last_seen else "never"
                        entries.append(f"**{rank}.** **{short_name}** - Lv.{level} ({xp:,} XP) - seen {seen}")
            except Exception as e:
                entries = [f"❌ Error loading level data: {e}"]

            total_pages = (total_count + limit - 1) // limit if total_count > 0 else 1

            embed.add_field(
                name=f"🏆 **Top Users by XP (Page {page}/{total_pages})**",
                value=("\n".join(entries) or "No users with XP yet")[:1024],
                inline=False
            )

            return embed, total_pages

        category_builders = [
            create_user_activity_embed,
            create_gaming_sessions_embed,
            create_listening_sessions_embed,
            create_active_gaming_embed,
            create_active_listening_embed,
            create_top_levels_embed,
        ]

        async def build_category_embed(category, page):
            """Build the embed for a log category and add the navigation footer"""
            embed, total_pages = await category_builders[category](page)
            embed.set_footer(
                text=f"👑 Page {page} of {total_pages} • Category: {category_names[category]} • Limit: {limit} entries",
                icon_url=ctx.author.display_avatar.url
            )
            return embed, total_pages

        # Create dropdown select menu for category navigation
        class LogsSelect(discord.ui.Select):
            def __init__(self):
                options = []
                for i, (category_name, category_data) in enumerate(log_categories.items()):
                    options.append(
                        discord.SelectOption(
                            label=category_name,
                            description=category_data["description"],
                            emoji=category_data["emoji"],
                            value=str(i)
                        )
                    )
                
                super().__init__(
                    placeholder="📋 Select a log category...",
                    min_values=1,
                    max_values=1,
                    options=options
                )
            
            async def callback(self, interaction):
                nonlocal current_category, current_page
                current_category = int(self.values[0])
                current_page = 1  # Reset to page 1 when changing category

                embed, total_pages = await build_category_embed(current_category, current_page)

                # Update view with new pagination buttons if needed
                view = LogsView(current_category, current_page, total_pages, self.view.admin_cog)
                await interaction.response.edit_message(embed=embed, view=view)
        
        # Create view with dropdown and navigation buttons
        class LogsView(OwnerOnlyView):
            def __init__(self, category, page, total_pages, admin_cog):
                super().__init__(ctx.author.id, timeout=300)  # Admin only, 5 minute timeout
                LogsView.latest = self  # Each navigation attaches a new view - only the newest one counts
                self.category = category
                self.page = page
                self.total_pages = total_pages
                self.admin_cog = admin_cog
                
                # Add dropdown first
                self.add_item(LogsSelect())
                
                # Add pagination buttons only if there's more than one page
                if self.total_pages > 1:
                    self.add_item(PreviousButton(self.page <= 1))
                    self.add_item(NextButton(self.page >= self.total_pages))
                
                self.add_item(OverviewButton(self.admin_cog))
                self.add_item(CloseButton())
            
            async def update_view(self, interaction, new_page=None):
                """Update the view with new state"""
                nonlocal current_page
                
                if new_page is not None:
                    current_page = new_page

                embed, total_pages = await build_category_embed(current_category, current_page)

                # Create new view with updated state
                view = LogsView(current_category, current_page, total_pages, self.admin_cog)
                await interaction.response.edit_message(embed=embed, view=view)
            
            async def on_timeout(self):
                # Remove the logs from the channel once the admin stops using them
                if LogsView.latest is self and getattr(LogsView, 'logs_message', None):
                    try:
                        await LogsView.logs_message.delete()
                    except discord.HTTPException:
                        pass
        
        # Button classes (same as before)
        class PreviousButton(discord.ui.Button):
            def __init__(self, disabled=False):
                super().__init__(label='◀️ Previous', style=discord.ButtonStyle.secondary, disabled=disabled)
            
            async def callback(self, interaction):
                if self.view.page > 1:
                    await self.view.update_view(interaction, new_page=self.view.page - 1)
        
        class NextButton(discord.ui.Button):
            def __init__(self, disabled=False):
                super().__init__(label='▶️ Next', style=discord.ButtonStyle.secondary, disabled=disabled)
            
            async def callback(self, interaction):
                if self.view.page < self.view.total_pages:
                    await self.view.update_view(interaction, new_page=self.view.page + 1)
        
        class OverviewButton(discord.ui.Button):
            def __init__(self, admin_cog):
                super().__init__(label='🏠 Overview', style=discord.ButtonStyle.primary)
                self.admin_cog = admin_cog
            
            async def callback(self, interaction):
                # Create overview embed
                overview_embed = discord.Embed(
                    title="📊 System Overview & Performance",
                    description="**Bot performance metrics and system statistics**",
                    color=discord.Color.gold()
                )
                
                # System performance
                import psutil
                
                cpu_percent = psutil.cpu_percent()
                memory = psutil.virtual_memory()
                memory_mb = memory.used // 1024 // 1024
                memory_percent = memory.percent
                
                uptime = 0.0
                if hasattr(self.admin_cog.bot, 'start_time') and self.admin_cog.bot.start_time:
                    uptime = time.time() - self.admin_cog.bot.start_time
                uptime_hours = int(uptime // 3600)
                uptime_mins = int((uptime % 3600) // 60)
                
                system_info = f"💻 **CPU Usage:** {cpu_percent:.1f}%\n🧠 **RAM Usage:** {memory_mb}MB ({memory_percent:.1f}%)\n📡 **Latency:** {round(self.admin_cog.bot.latency * 1000)}ms\n🕐 **Uptime:** {uptime_hours}h {uptime_mins}m"
                
                overview_embed.add_field(
                    name="⚙️ **System Performance**",
                    value=system_info,
                    inline=True
                )
                
                # Bot statistics with music
                active_listeners = len(getattr(self.admin_cog.bot, 'active_listening', {}))
                bot_info = f"🏠 **Servers:** {len(self.admin_cog.bot.guilds)}\n👥 **Total Users:** {len(self.admin_cog.bot.users)}\n🎮 **Active Gamers:** {len(getattr(self.admin_cog.bot, 'active_sessions', {}))}\n🎵 **Active Listeners:** {active_listeners}\n🤖 **Loaded Cogs:** {len(self.admin_cog.bot.cogs)}"
                
                overview_embed.add_field(
                    name="🤖 **Bot Statistics**",
                    value=bot_info,
                    inline=True
                )
                
                # Quick stats with music
                try:
                    async with aiosqlite.connect(self.admin_cog.db.db_path) as db:
                        async with db.execute("SELECT COUNT(*) FROM users") as cursor:
                            result = await cursor.fetchone()
                            total_tracked = result[0] if result else 0
                        
                        async with db.execute("SELECT COUNT(*) FROM game_sessions WHERE duration IS NOT NULL") as cursor:
                            result = await cursor.fetchone()
                            total_sessions = result[0] if result else 0
                        
                        async with db.execute("SELECT COUNT(*) FROM listening_sessions WHERE duration IS NOT NULL") as cursor:
                            result = await cursor.fetchone()
                            total_listening_sessions = result[0] if result else 0
                        
                        async with db.execute("SELECT SUM(total_messages) FROM users") as cursor:
                            result = await cursor.fetchone()
                            total_messages = result[0] if result and result[0] else 0
                    
                    quick_stats = f"📊 **{total_tracked}** tracked users\n💬 **{total_messages:,}** total messages\n🎲 **{total_sessions:,}** gaming sessions\n🎵 **{total_listening_sessions:,}** listening sessions\n⚡ **Live** data tracking"
                    
                    overview_embed.add_field(
                        name="📈 **Quick Statistics**",
                        value=quick_stats,
                        inline=True
                    )
                except Exception as e:
                    overview_embed.add_field(
                        name="📈 **Quick Statistics**",
                        value="❌ Error loading database stats",
                        inline=True
                    )
                
                # Categories summary
                categories_summary = ""
                for i, (category_name, category_data) in enumerate(log_categories.items()):
                    categories_summary += f"{category_data['emoji']} **{category_name}**\n• {category_data['description']}\n\n"
                
                overview_embed.add_field(
                    name="📋 **Available Log Categories**",
                    value=categories_summary.strip(),
                    inline=False
                )
                
                overview_embed.set_footer(text=f"👑 Use the dropdown menu to explore specific log categories • Limit: {limit} entries")
                if self.admin_cog.bot.user:
                    overview_embed.set_thumbnail(url=self.admin_cog.bot.user.avatar.url if self.admin_cog.bot.user.avatar else self.admin_cog.bot.user.default_avatar.url)
                
                await interaction.response.edit_message(embed=overview_embed, view=self.view)
        
        class CloseButton(discord.ui.Button):
            def __init__(self):
                super().__init__(label='❌ Close', style=discord.ButtonStyle.danger)
            
            async def callback(self, interaction):
                await interaction.response.edit_message(
                    embed=discord.Embed(
                        title="📋 Logs Menu Closed",
                        description="Use `!r logs` to open again.",
                        color=discord.Color.red()
                    ),
                    view=None
                )
                if self.view:
                    self.view.stop()
        
        # Create initial embed and view
        embed, total_pages = await build_category_embed(current_category, current_page)
        view = LogsView(current_category, current_page, total_pages, self)

        # Private reply ("Only you can see this") - removed when the menu times out
        logs_message = await ctx.send(embed=embed, view=view, ephemeral=True)
        LogsView.logs_message = logs_message

async def setup(bot):
    await bot.add_cog(RoxyAdmin(bot))
    print("👑 Roxy's Admin cog loaded with music support!")